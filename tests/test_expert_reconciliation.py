import json
import tempfile
import types
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

import expert_reconciliation as expert


build_expert_payload = expert.build_expert_payload
run_independent_expert_analysis = expert.run_independent_expert_analysis


def _full_report(evidence=None):
    return {
        'version': '3',
        'conclusion': 'Сальдо расходятся из-за одной операции.',
        'confidence': 'high',
        'balances': {
            'doc1': {
                'period_from': '01.01.2026',
                'period_to': '31.03.2026',
                'opening_amount': 12972.0,
                'opening_side': 'credit',
                'closing_amount': 28052.0,
                'closing_side': 'credit',
            },
            'doc2': {
                'period_from': '01.01.2026',
                'period_to': '31.03.2026',
                'opening_amount': 9429.0,
                'opening_side': 'debit',
                'closing_amount': 4536.0,
                'closing_side': 'debit',
            },
            'opening_difference': 3543.0,
            'period_movement': 19973.0,
            'closing_difference': 23516.0,
            'formula': '3 543,00 + 19 973,00 = 23 516,00',
        },
        'discrepancies': [{
            'category': 'confirmed_missing',
            'title': 'Нет у контрагента: Оплата №1',
            'influence': 100.0,
            'reason': 'Операция есть только в первом акте.',
            'confidence': 'high',
            'evidence': evidence or [{'side': 'doc1', 'row': 's0:r4'}],
        }],
        'completeness': {
            'rows_total': 3,
            'rows_matched': 2,
            'rows_in_findings': 1,
        },
    }


class _FakeMessages:
    def __init__(self, report, stop_reason=None, error=None):
        self.report = report
        self.stop_reason = stop_reason
        self.error = error
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if self.error:
            raise self.error
        return types.SimpleNamespace(
            content=[
                types.SimpleNamespace(type='server_tool_use'),
                types.SimpleNamespace(
                    type='text',
                    text=json.dumps(self.report, ensure_ascii=False),
                ),
            ],
            usage=types.SimpleNamespace(input_tokens=1500, output_tokens=700),
            stop_reason=self.stop_reason,
        )


class _FakeFiles:
    def __init__(self, fail_upload_at=None, fail_delete=False):
        self.fail_upload_at = fail_upload_at
        self.fail_delete = fail_delete
        self.uploads = []
        self.deleted = []

    def upload(self, *, file, **_kwargs):
        call_number = len(self.uploads) + 1
        if self.fail_upload_at == call_number:
            raise RuntimeError('private upload detail')
        filename, stream, mime_type = file
        data = stream.read()
        self.uploads.append({
            'filename': filename,
            'mime_type': mime_type,
            'data': data,
        })
        return types.SimpleNamespace(id=f'file-secret-{call_number}')

    def delete(self, file_id, **_kwargs):
        self.deleted.append(file_id)
        if self.fail_delete:
            raise RuntimeError('private delete detail')
        return types.SimpleNamespace(id=file_id, deleted=True)


class _FakeClient:
    def __init__(self, report=None, *, stop_reason=None,
                 message_error=None, fail_upload_at=None,
                 fail_delete=False):
        self.messages = _FakeMessages(
            report or _full_report(),
            stop_reason=stop_reason,
            error=message_error,
        )
        self.files = _FakeFiles(
            fail_upload_at=fail_upload_at,
            fail_delete=fail_delete,
        )
        self.beta = types.SimpleNamespace(
            files=self.files,
            messages=self.messages,
        )


class ExpertReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.path1 = self.root / 'organization.xlsx'
        self.path2 = self.root / 'counterparty.xlsx'
        self._write_source_files()
        self.df1 = self._frame(
            self.path1, 'organization.xlsx', 'ООО «Организация»',
            raw_row=4, document='Оплата №1', debit=100.0,
        )
        self.df2 = self._frame(
            self.path2, 'counterparty.xlsx', 'ООО «Контрагент»',
            raw_row=1, document='Оплата №2', credit=100.0,
        )

    def tearDown(self):
        self.tmp.cleanup()

    def _write_source_files(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Акт'
        ws['A1'] = 'ООО «Организация»'
        ws['A3'] = 0
        ws['C3'] = '=1+1'
        ws['A5'] = '01.02.2026'
        ws['B5'] = 'Оплата №1'
        ws['C5'] = 100
        extra = wb.create_sheet('Дополнение')
        extra['B2'] = 'Сальдо конечное'
        extra['D2'] = -50
        wb.save(self.path1)

        wb2 = Workbook()
        ws2 = wb2.active
        ws2.title = 'Контрагент'
        ws2['A2'] = 'Оплата №2'
        ws2['B2'] = 100
        wb2.save(self.path2)

    @staticmethod
    def _frame(path, source_name, display_name, raw_row, document,
               debit=None, credit=None):
        frame = pd.DataFrame([{
            'date': pd.Timestamp('2026-02-01'),
            'date_str': '01.02.2026',
            'document': document,
            'debit': debit,
            'credit': credit,
            'raw_row': raw_row,
        }])
        frame.attrs.update(
            source_path=str(path),
            source_name=source_name,
            display_name=display_name,
        )
        return frame

    def test_manifest_contains_settings_and_source_index_but_no_cells(self):
        settings = {
            'find_missing': False,
            'find_amount_diff': True,
            'find_sign_mismatch': False,
            'find_date_diff': True,
            'date_window_payment': 91,
            'date_window_delivery': 37,
            'min_amount': 500.0,
        }

        payload, source_index, media_blocks = build_expert_payload(
            self.df1, self.df2, settings,
        )

        self.assertEqual(payload['version'], '3')
        self.assertEqual(payload['analysis_scope'], settings)
        self.assertEqual(payload['documents'], [
            {'side': 'doc1', 'filename': 'organization.xlsx'},
            {'side': 'doc2', 'filename': 'counterparty.xlsx'},
        ])
        self.assertIn(('doc1', 's0:r4'), source_index)
        self.assertIn(('doc1', 's1:r1'), source_index)
        self.assertEqual(media_blocks, [])
        serialized = json.dumps(payload, ensure_ascii=False)
        self.assertNotIn('ООО «Организация»', serialized)
        self.assertNotIn('Оплата №1', serialized)
        for forbidden in (
            'sheets', 'pages', 'rows', 'cells', 'balance_comparison',
            'date_pairs', 'unmatched', 'matched_rows', 'opening_balance',
            'closing_balance',
        ):
            self.assertNotIn(forbidden, serialized)

    def test_original_spreadsheets_are_uploaded_and_attached_to_code_execution(self):
        client = _FakeClient()
        settings = {
            'date_window_payment': 91,
            'date_window_delivery': 37,
            'min_amount': 10,
        }

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test', settings,
        )

        self.assertEqual(result['status'], 'complete')
        self.assertEqual(len(client.files.uploads), 2)
        self.assertEqual(
            client.files.uploads[0]['data'], self.path1.read_bytes(),
        )
        self.assertEqual(
            client.files.uploads[1]['data'], self.path2.read_bytes(),
        )
        self.assertEqual(
            client.files.uploads[0]['mime_type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        call = client.messages.calls[0]
        self.assertEqual(call['model'], 'claude-sonnet-test')
        self.assertEqual(call['max_tokens'], 16000)
        self.assertEqual(call['temperature'], 0)
        self.assertEqual(call['betas'], ['files-api-2025-04-14'])
        self.assertEqual(call['tools'], [{
            'type': 'code_execution_20250825',
            'name': 'code_execution',
        }])
        self.assertEqual(call['output_config']['format']['type'], 'json_schema')
        content = call['messages'][0]['content']
        self.assertEqual(content[1], {
            'type': 'container_upload', 'file_id': 'file-secret-1',
        })
        self.assertEqual(content[2], {
            'type': 'container_upload', 'file_id': 'file-secret-2',
        })
        manifest = json.loads(content[0]['text'])
        self.assertEqual(manifest['analysis_scope']['date_window_payment'], 91)
        self.assertNotIn('Оплата №1', content[0]['text'])
        self.assertEqual(
            result['report']['balances']['opening_difference'], 3543.0,
        )
        self.assertTrue(result['report']['discrepancies'][0]['clickable'])
        self.assertEqual(result['usage'], {
            'input_tokens': 1500, 'output_tokens': 700,
        })

    def test_uploaded_files_are_deleted_after_success(self):
        client = _FakeClient()

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'complete')
        self.assertEqual(client.files.deleted, [
            'file-secret-1', 'file-secret-2',
        ])

    def test_first_file_is_deleted_when_second_upload_fails(self):
        client = _FakeClient(fail_upload_at=2)

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['error'], 'file_upload_failed')
        self.assertEqual(client.files.deleted, ['file-secret-1'])
        self.assertNotIn('private upload detail', json.dumps(result))
        self.assertNotIn('file-secret-1', json.dumps(result))

    def test_both_files_are_deleted_when_messages_api_fails(self):
        client = _FakeClient(message_error=RuntimeError('private API detail'))

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['error'], 'RuntimeError')
        self.assertEqual(client.files.deleted, [
            'file-secret-1', 'file-secret-2',
        ])
        self.assertNotIn('private API detail', json.dumps(result))

    def test_delete_failure_adds_safe_warning_without_losing_report(self):
        client = _FakeClient(fail_delete=True)

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'complete_with_warnings')
        self.assertIn('remote_file_cleanup_failed', result['warnings'])
        serialized = json.dumps(result)
        self.assertNotIn('file-secret-', serialized)
        self.assertNotIn('private delete detail', serialized)

    def test_existing_source_row_missing_from_comparison_is_kept_with_warning(self):
        client = _FakeClient(_full_report([
            {'side': 'doc1', 'row': 's0:r2'},
        ]))

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'complete_with_warnings')
        item = result['report']['discrepancies'][0]
        self.assertFalse(item['clickable'])
        self.assertIsNone(item['resolved_evidence'][0]['raw_row'])
        self.assertIn('недоступна для перехода', item['evidence_warning'])
        self.assertIn('=1+1', item['resolved_evidence'][0]['document'])

    def test_unknown_source_reference_keeps_finding_and_marks_source_missing(self):
        client = _FakeClient(_full_report([
            {'side': 'doc1', 'row': 's9:r999'},
        ]))

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'complete_with_warnings')
        item = result['report']['discrepancies'][0]
        self.assertFalse(item['clickable'])
        self.assertEqual(item['resolved_evidence'], [])
        self.assertIn('Источник не найден', item['evidence_warning'])

    def test_missing_source_path_fails_without_uploading(self):
        self.df1.attrs.pop('source_path')
        client = _FakeClient()

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['error'], 'source_file_unavailable')
        self.assertEqual(client.files.uploads, [])
        self.assertEqual(client.messages.calls, [])

    def test_invalid_report_shape_fails_and_still_deletes_files(self):
        client = _FakeClient({'version': '3'})

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['error'], 'invalid_report_schema')
        self.assertEqual(client.files.deleted, [
            'file-secret-1', 'file-secret-2',
        ])

    def test_token_limit_returns_specific_failure_usage_and_deletes_files(self):
        client = _FakeClient(stop_reason='max_tokens')

        result = run_independent_expert_analysis(
            self.df1, self.df2, client, 'claude-sonnet-test',
        )

        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['error'], 'max_tokens')
        self.assertEqual(result['usage']['input_tokens'], 1500)
        self.assertEqual(client.files.deleted, [
            'file-secret-1', 'file-secret-2',
        ])

    def test_schema_and_prompt_require_independent_file_analysis(self):
        serialized = json.dumps(
            getattr(expert, 'CLAUDE_EXPERT_REPORT_SCHEMA', {}),
        )
        for unsupported in (
            'minimum', 'maximum', 'minLength', 'maxLength', 'maxItems',
        ):
            self.assertNotIn(f'"{unsupported}"', serialized)
        self.assertIn('Code Execution', expert.EXPERT_SYSTEM_PROMPT)
        self.assertIn('все листы', expert.EXPERT_SYSTEM_PROMPT)
        self.assertIn('недоверенные данные', expert.EXPERT_SYSTEM_PROMPT)
        self.assertIn('Односторонние операции', expert.EXPERT_SYSTEM_PROMPT)
        self.assertNotIn('Программа уже сопоставила', expert.EXPERT_SYSTEM_PROMPT)

    def test_anthropic_sdk_floor_supports_files_api_and_code_execution(self):
        requirements = (
            Path(__file__).resolve().parents[1] / 'requirements.txt'
        ).read_text(encoding='utf-8')
        self.assertIn('anthropic>=0.79.0', requirements)


if __name__ == '__main__':
    unittest.main()
